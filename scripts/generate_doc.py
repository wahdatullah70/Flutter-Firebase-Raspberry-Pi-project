from docx import Document
from docx.shared import Pt

def add_heading(document, text, level=1):
    document.add_heading(text, level=level)


def add_paragraph(document, text, bold=False):
    p = document.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)


def add_bullet(document, items):
    for item in items:
        document.add_paragraph(item, style='List Bullet')


doc = Document()

doc.add_heading('How to Run the Solar App from a Windows PC on an Android Phone', 0)

add_paragraph(doc, 'This short guide is written for non-technical users. It explains the simple steps to get the Solar app running on an Android phone from a Windows computer. If you get stuck, share this document with someone who can help you follow the steps.', False)

add_heading(doc, 'What you will need', level=1)
add_bullet(doc, [
    'A Windows PC with internet access',
    'A USB cable to connect your phone to the PC',
    'The Android phone you want to use (charged and unlocked)',
    'A helpful person if you are not comfortable with computer steps'
])

add_heading(doc, 'Quick overview', level=1)
add_paragraph(doc, '- We prepare the phone to allow app installation, connect the phone to the PC, and then install and run the app so you can see it on the phone.', False)

add_heading(doc, 'Step 1 — Prepare your phone', level=1)
add_paragraph(doc, '1. Open Settings on your phone.', False)
add_paragraph(doc, '2. Find "About phone" and tap the Build Number or MIUI version seven times (this enables Developer options).', False)
add_paragraph(doc, '3. Go back to Settings → Additional settings → Developer options (or Developer settings).', False)
add_paragraph(doc, '4. Turn ON "USB debugging". If your phone is a Xiaomi device, also turn ON "Install via USB" if available. The phone may ask you to confirm or sign in to a Xiaomi account. Please follow the on-screen prompts.', False)

add_heading(doc, 'Step 2 — Connect your phone to the PC', level=1)
add_paragraph(doc, '1. Use the USB cable to connect the phone to the PC.', False)
add_paragraph(doc, '2. When the phone asks for permission to allow debugging from the PC, accept it (tap Allow).', False)
add_paragraph(doc, '3. Make sure the phone stays unlocked while the app installs.', False)

add_heading(doc, 'Step 3 — Install and run the app (simple method)', level=1)
add_paragraph(doc, 'If you have a friendly developer or someone who can run commands for you, they can use the following steps. If not, ask them to help and show them this document.', False)
add_paragraph(doc, 'A developer will do these three things:', False)
add_bullet(doc, [
    'Build the app on the PC',
    'Install the app on your phone',
    'Open the app so you can see the Solar dashboard on the phone screen'
])
add_paragraph(doc, 'If the phone shows an "Install" prompt or a security dialog during installation, please tap "Install" or "Allow" to continue.', False)

add_heading(doc, 'If the phone blocks installs (common on Xiaomi/MIUI phones)', level=1)
add_paragraph(doc, 'Xiaomi (MIUI) phones sometimes block apps installed from a computer for security. If you see a message like "Install canceled by user":', False)
add_bullet(doc, [
    'Make sure the phone screen is unlocked and you are watching the screen when the install runs.',
    'Enable "Install via USB" in Developer options (you may be asked to sign in to your Mi account).',
    'Try installing the app manually by opening the file manager on the phone and tapping the app file (the developer can copy the file into the phone).'
])

add_heading(doc, 'Optional: Seeing the phone screen on your PC', level=1)
add_paragraph(doc, 'There are tools that mirror your phone screen on your PC so you can watch and control it from the computer. One popular tool is "scrcpy". Ask a helper to install and run it if you want the convenience of viewing the phone on the PC screen.', False)

add_heading(doc, 'Troubleshooting tips', level=1)
add_bullet(doc, [
    'If the phone does not appear on the PC: check the USB cable and port and try again.',
    'If you see a message about drivers on Windows: ask your helper to install the phone manufacturer drivers, or install the Google USB driver via Android Studio.',
    'If the installation fails with security messages: accept any prompts on the phone and enable "Install via USB" if your phone has it.'
])

add_heading(doc, 'Help & next steps', level=1)
add_paragraph(doc, 'If you want, I can:', False)
add_bullet(doc, [
    'Retry installing the app while you watch and accept any prompts on the phone',
    'Walk you step-by-step through enabling "Install via USB" on a Xiaomi phone',
    'Provide a simple checklist you can share with a helper'
])

add_paragraph(doc, 'Location of this file from the project folder: windows_android_guide.docx', True)

# Save the document
output_path = '../windows_android_guide.docx'
doc.save(output_path)
print('Saved doc to:', output_path)
